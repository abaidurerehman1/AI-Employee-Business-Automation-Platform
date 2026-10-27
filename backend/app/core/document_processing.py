import io
import re
import unicodedata
import zipfile

from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

MAX_EXTRACTED_CHARACTERS = 2_000_000


def extract_text(content: bytes, extension: str) -> str:
    if extension == "txt":
        text = content.decode("utf-8-sig")
    elif extension == "pdf":
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported")
        if len(reader.pages) > 500:
            raise ValueError("PDF has too many pages")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif extension == "docx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Document expands beyond the allowed size")
        document = Document(io.BytesIO(content))
        text = "\n\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += "\n\n" + "\n".join(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
    else:
        raise ValueError("Unsupported document type")
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError("Extracted text exceeds the allowed size")
    text = unicodedata.normalize("NFKC", text).replace("\x00", "")
    text = re.sub(r"[\t ]+", " ", text)
    text = re.sub(r"\n[ \t]*\n(?:[ \t]*\n)+", "\n\n", text).strip()
    if not text:
        raise ValueError("No text could be extracted from this document")
    return text


def split_text(text: str) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1200, chunk_overlap=180, separators=["\n\n", "\n", ". ", " ", ""]
    )
    chunks = splitter.split_text(text)
    if len(chunks) > 2500:
        raise ValueError("Document produces too many chunks")
    return chunks
