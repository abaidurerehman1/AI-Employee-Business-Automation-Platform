import csv
import io
import uuid
import zipfile
import re
import unicodedata
from typing import Annotated, Literal

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, EmailStr, Field, ValidationError, field_validator
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError

from app.api.dependencies import DbSession, TenantContext, get_tenant_context, require_csrf, require_roles
from app.models.auth import Membership, MembershipRole, User
from app.models.crm import CRMActivity, Company, Contact, Deal, Lead

router = APIRouter(prefix="/crm", tags=["crm"])
Tenant = Annotated[TenantContext, Depends(get_tenant_context)]
Manager = Annotated[TenantContext, Depends(require_roles(MembershipRole.admin, MembershipRole.manager))]
Csrf = Annotated[User, Depends(require_csrf)]
Kind = Literal["companies", "contacts", "leads", "deals"]
MODELS = {"companies": Company, "contacts": Contact, "leads": Lead, "deals": Deal}


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)


class CompanyInput(StrictModel):
    name: str = Field(min_length=1, max_length=200)
    website: str | None = Field(default=None, max_length=500)


class ContactInput(StrictModel):
    name: str = Field(min_length=1, max_length=200)
    email: EmailStr | None = None
    phone: str | None = Field(default=None, max_length=40)
    company_id: uuid.UUID | None = None

    @field_validator("email", mode="before")
    @classmethod
    def normalize_email(cls, value):
        return value.strip().lower() if isinstance(value, str) and value.strip() else None

    @field_validator("phone", mode="before")
    @classmethod
    def normalize_phone(cls, value):
        if not value or not isinstance(value, str):
            return value
        normalized = re.sub(r"[\s().-]", "", value.strip())
        if not re.fullmatch(r"\+?[0-9]{3,25}", normalized):
            raise ValueError("Phone must contain 3 to 25 digits and an optional leading plus")
        return normalized


class LeadInput(StrictModel):
    title: str = Field(min_length=1, max_length=200)
    stage: Literal["new", "contacted", "qualified", "proposal", "won", "lost"] = "new"
    contact_id: uuid.UUID | None = None
    company_id: uuid.UUID | None = None
    owner_id: uuid.UUID | None = None


class DealInput(StrictModel):
    title: str = Field(min_length=1, max_length=200)
    stage: Literal["open", "negotiation", "won", "lost"] = "open"
    amount_minor: int = Field(default=0, ge=0, le=2_000_000_000)
    currency: str = Field(default="USD", pattern=r"^[A-Z]{3}$")
    lead_id: uuid.UUID | None = None
    owner_id: uuid.UUID | None = None


class ActivityInput(StrictModel):
    kind: Literal["call", "email", "note", "activity"]
    content: str = Field(min_length=1, max_length=10000)


INPUTS = {"companies": CompanyInput, "contacts": ContactInput, "leads": LeadInput, "deals": DealInput}


def serialize(record):
    return {column.name: getattr(record, column.name) for column in record.__table__.columns if column.name != "normalized_name"}


async def scoped(db, model, org_id, record_id):
    item = await db.scalar(select(model).where(model.id == record_id, model.organization_id == org_id))
    if not item:
        raise HTTPException(404, "CRM record not found")
    return item


async def validate_links(db, organization_id, data):
    for field, model in (("company_id", Company), ("contact_id", Contact), ("lead_id", Lead)):
        if data.get(field):
            await scoped(db, model, organization_id, data[field])
    if data.get("owner_id"):
        member = await db.scalar(select(Membership.id).where(
            Membership.organization_id == organization_id, Membership.user_id == data["owner_id"]))
        if not member:
            raise HTTPException(400, "Owner must belong to this organization")


def validate(kind, payload):
    try:
        data = INPUTS[kind].model_validate(payload).model_dump()
    except ValidationError as exc:
        raise HTTPException(422, exc.errors(include_url=False, include_context=False)) from exc
    if kind == "companies":
        data["normalized_name"] = " ".join(unicodedata.normalize("NFKC", data["name"]).casefold().split())
    return data


async def commit(db):
    try:
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(409, "A matching company or contact already exists") from exc


@router.post("/imports/preview")
async def preview_import(db: DbSession, context: Manager, _user: Csrf, file: Annotated[UploadFile, File()]):
    filename = (file.filename or "").lower()
    content = await file.read(5 * 1024 * 1024 + 1)
    await file.close()
    if not content or len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "File must be between 1 byte and 5 MB")
    try:
        if filename.endswith(".csv"):
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        elif filename.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError("Spreadsheet expands beyond limit")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                rows = []
                for row in book.active.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise ValueError("Formula cells are not supported")
                    rows.append([str(cell.value or "") for cell in row])
                    if len(rows) > 1001:
                        raise ValueError("At most 1000 rows are allowed")
            finally:
                book.close()
        else:
            raise HTTPException(415, "Upload a CSV or XLSX file")
        if not rows or len(rows) > 1001:
            raise ValueError("File is empty or has more than 1000 rows")
        headers = [str(value).strip().lower() for value in rows[0]]
        if len(headers) != len(set(headers)) or "name" not in headers or "email" not in headers:
            raise ValueError("Headers must include name and email without duplicates")
        result, seen = [], set()
        for number, row in enumerate(rows[1:], 2):
            raw = dict(zip(headers, row))
            if not any(str(v).strip() for v in raw.values()):
                continue
            try:
                contact = ContactInput(name=str(raw.get("name", "")), email=str(raw.get("email", "")), phone=str(raw.get("phone", "")) or None)
                if not contact.email:
                    raise ValueError("Email is required for import deduplication")
                duplicate = contact.email in seen or bool(await db.scalar(select(Contact.id).where(Contact.organization_id == context.organization.id, Contact.email == contact.email)))
                seen.add(contact.email)
                result.append({"row": number, "data": contact.model_dump(mode="json"), "status": "duplicate" if duplicate else "ready"})
            except (ValidationError, ValueError) as exc:
                result.append({"row": number, "data": raw, "status": "invalid", "error": str(exc)})
        return {"rows": result}
    except (ValueError, UnicodeError, zipfile.BadZipFile) as exc:
        raise HTTPException(400, str(exc)) from exc


class ImportInput(StrictModel):
    contacts: list[ContactInput] = Field(max_length=1000)


@router.post("/imports/contacts", status_code=201)
async def import_contacts(payload: ImportInput, db: DbSession, context: Manager, _user: Csrf):
    created = skipped = 0
    seen = set()
    for contact in payload.contacts:
        if not contact.email:
            raise HTTPException(422, "Email is required for import deduplication")
        await validate_links(db, context.organization.id, contact.model_dump())
        if contact.email in seen or await db.scalar(select(Contact.id).where(Contact.organization_id == context.organization.id, Contact.email == contact.email)):
            skipped += 1
            continue
        seen.add(contact.email)
        db.add(Contact(organization_id=context.organization.id, **contact.model_dump()))
        created += 1
    await commit(db)
    return {"created": created, "skipped": skipped}


@router.get("/{kind}")
async def list_records(kind: Kind, db: DbSession, context: Tenant, q: str = "", stage: str = "", owner_id: uuid.UUID | None = None, page: int = 1, page_size: int = 20):
    if len(q) > 200 or page < 1 or not 1 <= page_size <= 100:
        raise HTTPException(422, "Invalid pagination or search")
    model = MODELS[kind]
    query = select(model).where(model.organization_id == context.organization.id)
    if q.strip():
        field = model.name if kind in {"companies", "contacts"} else model.title
        term = q.strip().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        search = field.ilike(f"%{term}%", escape="\\")
        if kind == "contacts":
            search = or_(search, Contact.email.ilike(f"%{term}%", escape="\\"), Contact.phone.ilike(f"%{term}%", escape="\\"))
        query = query.where(search)
    if stage and hasattr(model, "stage"):
        query = query.where(model.stage == stage)
    if owner_id and hasattr(model, "owner_id"):
        query = query.where(model.owner_id == owner_id)
    total = await db.scalar(select(func.count()).select_from(query.subquery()))
    rows = (await db.scalars(query.order_by(model.created_at.desc(), model.id.desc()).offset((page - 1) * page_size).limit(page_size))).all()
    return {"items": [serialize(item) for item in rows], "total": total, "page": page, "page_size": page_size}


@router.post("/{kind}", status_code=201)
async def create_record(kind: Kind, payload: dict, db: DbSession, context: Manager, _user: Csrf):
    data = validate(kind, payload)
    await validate_links(db, context.organization.id, data)
    record = MODELS[kind](organization_id=context.organization.id, **data)
    db.add(record)
    await commit(db)
    await db.refresh(record)
    return serialize(record)


@router.get("/{kind}/{record_id}")
async def get_record(kind: Kind, record_id: uuid.UUID, db: DbSession, context: Tenant):
    return serialize(await scoped(db, MODELS[kind], context.organization.id, record_id))


@router.put("/{kind}/{record_id}")
async def update_record(kind: Kind, record_id: uuid.UUID, payload: dict, db: DbSession, context: Manager, _user: Csrf):
    record = await scoped(db, MODELS[kind], context.organization.id, record_id)
    data = validate(kind, payload)
    await validate_links(db, context.organization.id, data)
    for name, value in data.items():
        setattr(record, name, value)
    await commit(db)
    await db.refresh(record)
    return serialize(record)


@router.get("/leads/{lead_id}/activities")
async def list_activities(lead_id: uuid.UUID, db: DbSession, context: Tenant):
    await scoped(db, Lead, context.organization.id, lead_id)
    rows = (await db.scalars(select(CRMActivity).where(CRMActivity.organization_id == context.organization.id, CRMActivity.lead_id == lead_id).order_by(CRMActivity.created_at.desc()).limit(100))).all()
    return [serialize(row) for row in rows]


@router.post("/leads/{lead_id}/activities", status_code=201)
async def create_activity(lead_id: uuid.UUID, payload: ActivityInput, db: DbSession, context: Manager, _user: Csrf):
    await scoped(db, Lead, context.organization.id, lead_id)
    activity = CRMActivity(organization_id=context.organization.id, lead_id=lead_id, created_by=context.user.id, **payload.model_dump())
    db.add(activity)
    await db.commit()
    await db.refresh(activity)
    return serialize(activity)
